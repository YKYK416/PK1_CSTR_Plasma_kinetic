"""Export the submission Markdown manuscript to an editable A4 Word document.

This intentionally lightweight exporter preserves headings, paragraphs, tables,
embedded PNG figures, captions, and references without changing the source
Markdown manuscript or the figure assets.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "core" / "docs" / "manuscript" / "manuscript_submission_en_v1.md"
OUTPUT = ROOT / "core" / "docs" / "manuscript" / "manuscript_submission_en_v1.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    width = OxmlElement("w:tcW")
    width.set(qn("w:w"), str(int(width_cm * 567)))
    width.set(qn("w:type"), "dxa")
    tc_pr.append(width)


def add_page_number(paragraph) -> None:
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def tex_to_text(text: str) -> str:
    """Render the manuscript's limited LaTeX notation as readable editable text."""
    text = text.replace("\\[", "").replace("\\]", "")
    text = text.replace("\\(", "").replace("\\)", "")
    substitutions = {
        r"\\times": "×",
        r"\\leq": "≤",
        r"\\geq": "≥",
        r"\\approx": "≈",
        r"\\rightarrow": "→",
        r"\\sum": "Σ",
        r"\\tau": "τ",
        r"\\nu": "ν",
        r"\\mathrm": "",
        r"\\mathrm": "",
        r"\\text": "",
        r"\\,": " ",
    }
    for old, new in substitutions.items():
        text = text.replace(old, new)
    text = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", text)
    text = re.sub(r"\\([A-Za-z]+)\{([^{}]+)\}", r"\2", text)
    text = re.sub(r"\^\{([^{}]+)\}", r"^\1", text)
    text = re.sub(r"_\{([^{}]+)\}", r"_\1", text)
    text = text.replace("\\", "")
    return text


def clean_text(text: str) -> str:
    text = tex_to_text(text)
    text = text.replace("`", "")
    text = text.replace("**", "")
    text = text.replace("*", "")
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    return re.sub(r"\s+", " ", text).strip()


def configure_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in (("Title", 16, "17365D"), ("Heading 1", 13, "17365D"), ("Heading 2", 11.5, "1F4E79")):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(13 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    if "Figure Caption" not in doc.styles:
        caption = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Times New Roman"
        caption.font.size = Pt(9)
        caption.font.italic = True
        caption.paragraph_format.space_after = Pt(8)
        caption.paragraph_format.line_spacing = 1.0

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.style = doc.styles["Normal"]
    add_page_number(footer)
    return doc


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    usable_cm = 21.0 - 2.4 - 2.4
    widths = [usable_cm / len(rows[0])] * len(rows[0])
    for r_index, values in enumerate(rows):
        cells = table.add_row().cells
        for c_index, value in enumerate(values):
            cell = cells[c_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[c_index])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(clean_text(value))
            run.font.name = "Times New Roman"
            run.font.size = Pt(8.5)
            if r_index == 0:
                run.bold = True
                set_cell_shading(cell, "D9EAF7")
    doc.add_paragraph()


def add_figure(doc: Document, source: Path, alt: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(source), width=Inches(6.35))
    paragraph.paragraph_format.space_after = Pt(3)
    note = doc.add_paragraph(style="Figure Caption")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(clean_text(alt))


def export(output: Path = OUTPUT) -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = configure_document()
    index = 0
    seen_title = False

    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            image = (SOURCE.parent / image_match.group(2)).resolve()
            if not image.exists():
                raise FileNotFoundError(image)
            add_figure(doc, image, image_match.group(1))
            index += 1
            continue

        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(clean_text(line[2:]))
            seen_title = True
            index += 1
            continue

        if line.startswith("## "):
            heading = clean_text(line[3:])
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.add_run(heading)
            index += 1
            continue

        if line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.add_run(clean_text(line[4:]))
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?", lines[index + 1]):
            rows = []
            header = [cell.strip() for cell in line.strip().strip("|").split("|")]
            rows.append(header)
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(doc, rows)
            continue

        style = "Normal"
        if line.startswith("**Figure ") or line.startswith("**Table "):
            style = "Figure Caption"
        paragraph = doc.add_paragraph(style=style)
        if line.startswith("**Authors:") or line.startswith("**Affiliations:") or line.startswith("**Corresponding author:") or line.startswith("**Manuscript status:"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(clean_text(line))
        index += 1

    doc.core_properties.title = "Electronically Excited H2 Sustains Direct NH Entry while Downstream Turnover Controls NH3 in a Plasma CSTR"
    doc.core_properties.subject = "English manuscript submission draft, version 1.7 (five-module Methods)"
    doc.core_properties.comments = "Generated from manuscript_submission_en_v1.md; version 1.7 five-module Methods; Figures 1–26 and graphical abstract embedded."
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export the English manuscript Markdown to an editable DOCX file.")
    parser.add_argument("--output", type=Path, default=OUTPUT, help="DOCX path to create; defaults to the current submission draft.")
    args = parser.parse_args()
    export(args.output)
